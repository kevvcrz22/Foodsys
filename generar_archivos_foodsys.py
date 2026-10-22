import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import unicodedata
import re
import os

BASE_DIR = r"c:\Foodsys\Foodsys"

def normalizar_programa(val):
    if pd.isna(val) or val is None:
        return "PROGRAMA EN FORMACION SENA"
    s = str(val).strip().upper()
    s = s.replace("Ñ", "__ENIE__")
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    s = s.replace("__ENIE__", "Ñ")
    s = re.sub(r"^[.,;:_\-/\\]+|[.,;:_\-/\\]+$", "", s).strip()
    s = re.sub(r"\s+", " ", s)
    return s if s else "PROGRAMA EN FORMACION SENA"

def clean_doc(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if "e" in s.lower():
        try:
            return str(int(float(s)))
        except Exception:
            pass
    s = s.split(".")[0]
    return re.sub(r"\D", "", s)

def clean_phone(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if "e" in s.lower():
        try:
            return str(int(float(s)))
        except Exception:
            pass
    s = s.split(".")[0]
    digits = re.sub(r"\D", "", s)
    return digits if len(digits) >= 7 else ""

def clean_name(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip().upper()
    return re.sub(r"\s+", " ", s)

def is_active_state(s):
    s = str(s).upper().strip()
    return any(a in s for a in ["FORMACION", "INDUCCION", "VISITA", "CONDICIONADO", "PASANT"])

print("[1/5] Consolidando fichas activas...")
fichas_dict = {}

# 1. Leer fichas actuales si existen
path_fichas = os.path.join(BASE_DIR, "Fichas_Activas_Foodsys.xlsx")
if os.path.exists(path_fichas):
    df_existing_f = pd.read_excel(path_fichas)
    for _, row in df_existing_f.iterrows():
        num = int(row["Num_Ficha"])
        fichas_dict[num] = {
            "Num_Ficha": num,
            "Nombre del Programa": normalizar_programa(row["Nombre del Programa"]),
            "Nivel de Formacion": str(row["Nivel de Formacion"]).strip().upper(),
            "FecIniLec_Ficha": str(row["FecIniLec_Ficha"])[:10],
            "FecFinLec_Ficha": str(row["FecFinLec_Ficha"])[:10],
            "FecIniPra_Ficha": str(row["FecIniPra_Ficha"])[:10],
            "FecFinPra_Ficha": str(row["FecFinPra_Ficha"])[:10],
            "Estado": "Activo"
        }
print(f"  Fichas previas cargadas: {len(fichas_dict)}")

# 2. Leer Base_NO MODIFICAR_2026_SEPTIEMBRE.xlsm
path_base = os.path.join(BASE_DIR, "Base_NO MODIFICAR_2026_SEPTIEMBRE.xlsm")
wb_b = openpyxl.load_workbook(path_base, data_only=True)
ws_b = wb_b["Base_Datos_IQR"]
data_b = list(ws_b.iter_rows(values_only=True))
headers_b = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(data_b[3])]
df_b = pd.DataFrame(data_b[4:], columns=headers_b)
df_b = df_b[df_b["Documento"].notnull() & (df_b["Documento"] != "")].copy()
wb_b.close()

# Añadir fichas de aprendices activos de la Base
for _, row in df_b[df_b["ESTADO"].apply(is_active_state)].iterrows():
    raw_f = row["Ficha"]
    if pd.isna(raw_f):
        continue
    try:
        f_num = int(float(raw_f))
    except Exception:
        continue
    if f_num not in fichas_dict:
        prog = normalizar_programa(row.get("Programa de Formación", ""))
        nivel = str(row.get("Nivel de Formación", "TECNOLOGO")).strip().upper() or "TECNOLOGO"
        fichas_dict[f_num] = {
            "Num_Ficha": f_num,
            "Nombre del Programa": prog,
            "Nivel de Formacion": nivel,
            "FecIniLec_Ficha": "2024-01-15",
            "FecFinLec_Ficha": "2026-07-14",
            "FecIniPra_Ficha": "2026-07-15",
            "FecFinPra_Ficha": "2027-01-14",
            "Estado": "Activo"
        }

print(f"  Total fichas consolidadas activas: {len(fichas_dict)}")

# 3. Consolidar Aprendices Activos
print("[2/5] Consolidando aprendices activos...")
aprendices_dict = {}

path_aprendices = os.path.join(BASE_DIR, "Aprendices_Activos_Foodsys.xlsx")
if os.path.exists(path_aprendices):
    df_existing_a = pd.read_excel(path_aprendices)
    for _, row in df_existing_a.iterrows():
        doc = clean_doc(row["Documento"])
        if not doc:
            continue
        f_num = int(row["Ficha"]) if pd.notna(row["Ficha"]) else None
        cen_con = "Si" if str(row.get("Centro de Convivencia", "")).strip().lower() in ["si", "sí", "1", "true"] else "No"
        rol = "Aprendiz Interno" if cen_con == "Si" else "Aprendiz Externo"
        cor = str(row.get("Correo Electrónico", "")).strip().lower()
        if cor == "nan": cor = ""
        tel = clean_phone(row.get("Teléfono", ""))
        gen = str(row.get("Género", "M")).strip().upper()[:1]
        if gen not in ["M", "F"]: gen = ""
        aprendices_dict[doc] = {
            "Tipo de Documento": str(row.get("Tipo de Documento", "CC")).strip().upper() or "CC",
            "Documento": int(doc),
            "Nombres": clean_name(row.get("Nombres", "")),
            "Apellidos": clean_name(row.get("Apellidos", "")),
            "Ficha": f_num,
            "Centro de Convivencia": cen_con,
            "Correo Electrónico": cor,
            "Teléfono": tel,
            "Género": gen,
            "Rol": rol,
            "Estado": "En Formacion"
        }

# Incorporar y enriquecer desde Base_Datos_IQR
for _, row in df_b[df_b["ESTADO"].apply(is_active_state)].iterrows():
    doc = clean_doc(row["Documento"])
    if not doc:
        continue
    raw_tel = clean_phone(row.get("TELÉFOSI DE CONTACTO", ""))
    raw_cor = str(row.get("CORREO ELECTRÓNICO ", "")).strip().lower()
    if raw_cor == "nan": raw_cor = ""

    if doc in aprendices_dict:
        if not aprendices_dict[doc]["Teléfono"] and raw_tel:
            aprendices_dict[doc]["Teléfono"] = raw_tel
        if not aprendices_dict[doc]["Correo Electrónico"] and raw_cor:
            aprendices_dict[doc]["Correo Electrónico"] = raw_cor
    else:
        raw_f = row["Ficha"]
        f_num = int(float(raw_f)) if pd.notna(raw_f) else None
        cen_con_raw = str(row.get("CENTRO DE CONVIVENCIA", "NO")).strip().upper()
        cen_con = "Si" if cen_con_raw in ["SI", "SÍ", "1", "TRUE"] else "No"
        rol = "Aprendiz Interno" if cen_con == "Si" else "Aprendiz Externo"
        gen_raw = str(row.get("Genero", "")).strip().upper()
        gen = "M" if gen_raw.startswith("M") else ("F" if gen_raw.startswith("F") else "")

        aprendices_dict[doc] = {
            "Tipo de Documento": str(row.get("Tipo de Documento", "CC")).strip().upper() or "CC",
            "Documento": int(doc),
            "Nombres": clean_name(row.get("SImbres", "")),
            "Apellidos": clean_name(row.get("Apellidos", "")),
            "Ficha": f_num,
            "Centro de Convivencia": cen_con,
            "Correo Electrónico": raw_cor,
            "Teléfono": raw_tel,
            "Género": gen,
            "Rol": rol,
            "Estado": "En Formacion"
        }

print(f"  Total aprendices consolidados activos: {len(aprendices_dict)}")

# ESTILOS PROFESIONALES
header_fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_fill_green = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
header_font_white = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
data_font = Font(name="Segoe UI", size=10)
bold_font = Font(name="Segoe UI", size=10, bold=True)
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")
border_thin = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1")
)
zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

# 4. CONSTRUCCIÓN DE EXCEL 1: FICHAS ACTIVAS
print("[3/5] Generando Fichas_Activas_Foodsys.xlsx...")
wb1 = openpyxl.Workbook()
ws1 = wb1.active
ws1.title = "Fichas_Activas"

cols_f = [
    "Num_Ficha", "Nombre del Programa", "Nivel de Formacion",
    "FecIniLec_Ficha", "FecFinLec_Ficha", "FecIniPra_Ficha", "FecFinPra_Ficha", "Estado"
]
ws1.append(cols_f)
for c_idx in range(1, len(cols_f) + 1):
    cell = ws1.cell(row=1, column=c_idx)
    cell.fill = header_fill_blue
    cell.font = header_font_white
    cell.alignment = center_align

sorted_fichas = sorted(fichas_dict.values(), key=lambda x: x["Num_Ficha"])
for r_idx, f in enumerate(sorted_fichas, start=2):
    row_data = [
        f["Num_Ficha"],
        f["Nombre del Programa"],
        f["Nivel de Formacion"],
        f["FecIniLec_Ficha"],
        f["FecFinLec_Ficha"],
        f["FecIniPra_Ficha"],
        f["FecFinPra_Ficha"],
        f["Estado"]
    ]
    ws1.append(row_data)
    for c_idx in range(1, len(cols_f) + 1):
        cell = ws1.cell(row=r_idx, column=c_idx)
        cell.font = data_font
        cell.border = border_thin
        if c_idx in [1, 3, 4, 5, 6, 7, 8]:
            cell.alignment = center_align
        else:
            cell.alignment = left_align
        if r_idx % 2 == 1:
            cell.fill = zebra_fill

for col in ws1.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 14)

# Hoja 2 Guia Fichas
ws1_guia = wb1.create_sheet("GUIA_USUARIO_FICHAS")
guia_fichas = [
    ["SISTEMA FOODSYS - GUIA OFICIAL PARA IMPORTAR FICHAS", "", "", ""],
    ["Centro Agropecuario La Granja - SENA Regional Tolima", "", "", ""],
    ["", "", "", ""],
    ["1. ESTRUCTURA Y SIGNIFICADO DE LAS COLUMNAS", "", "", ""],
    ["Columna", "Tipo de Dato", "Obligatorio", "Descripcion y Reglas de Negocio"],
    ["Num_Ficha", "Numerico", "SI (Obligatorio)", "Numero de la ficha SENA (Ej: 3404817)."],
    ["Nombre del Programa", "Texto", "SI (Obligatorio)", "Nombre del programa. Si no existe en Foodsys, el sistema lo crea automaticamente."],
    ["Nivel de Formacion", "Texto", "SI (Obligatorio)", "TECNOLOGO, TECNICO, OPERARIO, AUXILIAR o CURSO ESPECIAL."],
    ["FecIniLec_Ficha", "Fecha (AAAA-MM-DD)", "Opcional", "Fecha de inicio de etapa lectiva."],
    ["FecFinLec_Ficha", "Fecha (AAAA-MM-DD)", "Recomendado", "Fin de lectiva. Al expirar, los aprendices se inactivan de forma automatica."],
    ["FecIniPra_Ficha", "Fecha (AAAA-MM-DD)", "Opcional", "Fecha de inicio de etapa practica (-6 meses de la terminacion)."],
    ["FecFinPra_Ficha", "Fecha (AAAA-MM-DD)", "Recomendado", "Fecha oficial de terminacion de la ficha."],
    ["Estado", "Texto", "SI (Obligatorio)", "Estado de la ficha: Activo."],
    ["", "", "", ""],
    ["2. PASOS PARA SUBIR ESTE ARCHIVO A FOODSYS (PARA PERSONAS NO PROGRAMADORAS)", "", "", ""],
    ["Paso 1", "Inicia sesion en Foodsys con tu cuenta de Administrador."],
    ["Paso 2", "En el menu lateral izquierdo, abre la seccion 'Fichas'."],
    ["Paso 3", "Haz clic en el boton azul 'Importar Fichas'."],
    ["Paso 4", "Arrastra o selecciona este archivo (Fichas_Activas_Foodsys.xlsx) o el reporte institucional PE-04_1.xlsx."],
    ["Paso 5", "El sistema detecta las fechas y programas automaticamente."],
    ["Paso 6", "Haz clic en 'Importar Fichas Seleccionadas'. ¡Listo!"],
    ["", "", "", ""],
    ["3. PREGUNTAS FRECUENTES", "", "", ""],
    ["¿Puedo subir el PE-04 directamente?", "SI. Foodsys lee automaticamente el reporte oficial PE-04 del SENA."],
    ["¿Que pasa si una ficha ya existe?", "El sistema la actualiza sin generar duplicados."],
    ["¿Como se calculan las fechas?", "Foodsys calcula automaticamente 6 meses para la practica si vienen vacias."]
]
for row in guia_fichas:
    ws1_guia.append(row)

ws1_guia["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
ws1_guia["A2"].font = Font(name="Segoe UI", size=11, italic=True, color="475569")
ws1_guia["A4"].font = Font(name="Segoe UI", size=12, bold=True, color="1E3A8A")
ws1_guia["A14"].font = Font(name="Segoe UI", size=12, bold=True, color="1E3A8A")
ws1_guia["A22"].font = Font(name="Segoe UI", size=12, bold=True, color="1E3A8A")

for c in range(1, 5):
    cell = ws1_guia.cell(row=5, column=c)
    cell.fill = header_fill_blue
    cell.font = header_font_white
    cell.alignment = center_align

ws1_guia.column_dimensions["A"].width = 24
ws1_guia.column_dimensions["B"].width = 32
ws1_guia.column_dimensions["C"].width = 20
ws1_guia.column_dimensions["D"].width = 75

wb1.save(path_fichas)
print(f"  Archivo guardado: {path_fichas}")

# 5. CONSTRUCCIÓN DE EXCEL 2: APRENDICES ACTIVOS
print("[4/5] Generando Aprendices_Activos_Foodsys.xlsx...")
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "Aprendices_Activos"

cols_a = [
    "Tipo de Documento", "Documento", "Nombres", "Apellidos", "Ficha",
    "Centro de Convivencia", "Correo Electrónico", "Teléfono", "Género", "Rol", "Estado"
]
ws2.append(cols_a)
for c_idx in range(1, len(cols_a) + 1):
    cell = ws2.cell(row=1, column=c_idx)
    cell.fill = header_fill_green
    cell.font = header_font_white
    cell.alignment = center_align

sorted_aprendices = sorted(aprendices_dict.values(), key=lambda x: (x["Ficha"] or 0, x["Apellidos"], x["Nombres"]))
for r_idx, a in enumerate(sorted_aprendices, start=2):
    row_data = [
        a["Tipo de Documento"],
        a["Documento"],
        a["Nombres"],
        a["Apellidos"],
        a["Ficha"],
        a["Centro de Convivencia"],
        a["Correo Electrónico"],
        a["Teléfono"],
        a["Género"],
        a["Rol"],
        a["Estado"]
    ]
    ws2.append(row_data)
    for c_idx in range(1, len(cols_a) + 1):
        cell = ws2.cell(row=r_idx, column=c_idx)
        cell.font = data_font
        cell.border = border_thin
        if c_idx in [1, 2, 5, 6, 8, 9, 10, 11]:
            cell.alignment = center_align
        else:
            cell.alignment = left_align
        if r_idx % 2 == 1:
            cell.fill = zebra_fill

for col in ws2.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

# Hoja 2 Guia Aprendices
ws2_guia = wb2.create_sheet("GUIA_USUARIO_APRENDICES")
guia_aprendices = [
    ["SISTEMA FOODSYS - GUIA OFICIAL PARA IMPORTAR APRENDICES", "", "", ""],
    ["Centro Agropecuario La Granja - SENA Regional Tolima", "", "", ""],
    ["", "", "", ""],
    ["1. ESTRUCTURA Y SIGNIFICADO DE LAS COLUMNAS", "", "", ""],
    ["Columna", "Tipo de Dato", "Obligatorio", "Descripcion y Reglas de Negocio"],
    ["Tipo de Documento", "Texto (CC, TI, CE, PPT)", "SI (Obligatorio)", "Tipo de identificacion oficial del aprendiz."],
    ["Documento", "Numerico", "SI (Obligatorio)", "Numero de identificacion. Tambien es la contrasena inicial del aprendiz."],
    ["Nombres", "Texto (Mayusculas)", "SI (Obligatorio)", "Nombres del aprendiz."],
    ["Apellidos", "Texto (Mayusculas)", "SI (Obligatorio)", "Apellidos del aprendiz."],
    ["Ficha", "Numerico", "SI (Recomendado)", "Numero de ficha asignada."],
    ["Centro de Convivencia", "Texto (Si o No)", "SI (Obligatorio)", "'Si' para internos en cabanas (3 comidas). 'No' para externos (Almuerzo)."],
    ["Correo Electrónico", "Texto / Email", "Opcional", "Correo para restablecer contrasena y notificaciones."],
    ["Teléfono", "Numerico (10 digitos)", "Opcional", "Celular de contacto del aprendiz."],
    ["Género", "Texto (M o F)", "Opcional", "Genero del aprendiz."],
    ["Rol", "Texto", "SI (Obligatorio)", "'Aprendiz Interno' o 'Aprendiz Externo'."],
    ["Estado", "Texto", "SI (Obligatorio)", "Debe ser 'En Formacion' para poder hacer reservas."],
    ["", "", "", ""],
    ["2. PASOS PARA SUBIR ESTE ARCHIVO A FOODSYS (PARA PERSONAS NO PROGRAMADORAS)", "", "", ""],
    ["Paso 1", "Inicia sesion en Foodsys (Administrador, Coordinador o Bienestar)."],
    ["Paso 2", "En el menu lateral izquierdo, abre la seccion 'Aprendices' o 'Usuarios'."],
    ["Paso 3", "Haz clic en el boton azul 'Importar Excel'."],
    ["Paso 4", "Arrastra o selecciona este archivo (Aprendices_Activos_Foodsys.xlsx) o la base Base_NO MODIFICAR_2026_SEPTIEMBRE.xlsm."],
    ["Paso 5", "El sistema mostrara la vista previa con busqueda y paginacion rapida."],
    ["Paso 6", "Haz clic en 'Importar Seleccionados'. ¡Listo!"],
    ["", "", "", ""],
    ["3. REGLAS DE NEGOCIO IMPORTANTES", "", "", ""],
    ["Aprendiz Interno (Convivencia: Si)", "Tiene derecho a reservar Desayuno, Almuerzo y Cena."],
    ["Aprendiz Externo (Convivencia: No)", "Tiene derecho unicamente a reservar Almuerzo."],
    ["Contrasena Inicial", "El aprendiz ingresa con su numero de documento como usuario y contrasena."],
    ["Regla de Inasistencia", "A las 3 inasistencias en la misma semana, el sistema sanciona de forma automatica."]
]
for row in guia_aprendices:
    ws2_guia.append(row)

ws2_guia["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="065F46")
ws2_guia["A2"].font = Font(name="Segoe UI", size=11, italic=True, color="475569")
ws2_guia["A4"].font = Font(name="Segoe UI", size=12, bold=True, color="065F46")
ws2_guia["A17"].font = Font(name="Segoe UI", size=12, bold=True, color="065F46")
ws2_guia["A25"].font = Font(name="Segoe UI", size=12, bold=True, color="065F46")

for c in range(1, 5):
    cell = ws2_guia.cell(row=5, column=c)
    cell.fill = header_fill_green
    cell.font = header_font_white
    cell.alignment = center_align

ws2_guia.column_dimensions["A"].width = 24
ws2_guia.column_dimensions["B"].width = 32
ws2_guia.column_dimensions["C"].width = 20
ws2_guia.column_dimensions["D"].width = 75

wb2.save(path_aprendices)
print(f"  Archivo guardado: {path_aprendices}")
print("[5/5] ¡Archivos generados exitosamente con sus hojas de guia!")
